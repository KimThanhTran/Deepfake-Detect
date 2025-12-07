import csv
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:
    raise SystemExit("Please install openpyxl: pip install openpyxl")

ROOT = Path('.')
CSV_TOTAL = ROOT / 'TOTAL_RESULTS.csv'
CSV_DIFF_FOR = ROOT / 'results' / 'evaluation_details_diffusionforensics.csv'
CSV_DIFF_1K = ROOT / 'results' / 'evaluation_details_diffusion1k.csv'
XLSX_OUT = ROOT / 'TOTAL_RESULT.xlsx'

def read_csv_rows(path):
    if not path.exists():
        return [], []
    with path.open('r', encoding='utf-8', newline='') as f:
        reader = csv.reader(f)
        rows = list(reader)
        if not rows:
            return [], []
        header = rows[0]
        data = rows[1:]
        return header, data

def write_sheet(ws, title, header, data):
    ws.title = title
    if header:
        ws.append(header)
    for row in data:
        ws.append(row)
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = cell.value
            if val is None:
                continue
            val_str = str(val)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

def main():
    wb = Workbook()
    # Default sheet
    ws0 = wb.active
    ws0.title = 'README'
    ws0.append(["NPR Deepfake Detection — Tổng hợp kết quả"])
    ws0.append(["Xem các sheet: TOTAL_RESULTS, DiffusionForensics, Diffusion1kStep."])

    # TOTAL_RESULTS
    hdr, data = read_csv_rows(CSV_TOTAL)
    ws_total = wb.create_sheet('TOTAL_RESULTS')
    write_sheet(ws_total, 'TOTAL_RESULTS', hdr, data)

    # DiffusionForensics details
    hdr_df, data_df = read_csv_rows(CSV_DIFF_FOR)
    ws_df = wb.create_sheet('DiffusionForensics')
    write_sheet(ws_df, 'DiffusionForensics', hdr_df, data_df)

    # Diffusion1kStep details
    hdr_d1k, data_d1k = read_csv_rows(CSV_DIFF_1K)
    ws_d1k = wb.create_sheet('Diffusion1kStep')
    write_sheet(ws_d1k, 'Diffusion1kStep', hdr_d1k, data_d1k)

    wb.save(XLSX_OUT)
    print(f"Saved Excel workbook to {XLSX_OUT}")

if __name__ == '__main__':
    main()
